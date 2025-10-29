from __future__ import annotations

import json
import os
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Tuple, Dict, List, Any

from sqlalchemy import select, and_
from sqlalchemy.orm import Session

from config import ENABLE_SURCHARGES, COMPANY_LOGO_PATH
from .rules_service import calc_surcharges_for_entry, surcharge_amounts_from_minutes
from .wage_service import get_hour_rates as _wage_per_hour
from ..models import WorkEntry, UserProfile, PayrollRun, User
from ..utils.formatting import fmt_money


def _effective_hour_rates(session, user_id) -> tuple[float,float]:
    from core.models import UserProfile
    from config import DEFAULT_WAGE_PRESETS as _W
    p = session.query(UserProfile).filter(UserProfile.user_id==user_id).one_or_none()
    def _v(name, default=0.0):
        return (getattr(p, name) if (p and getattr(p, name) is not None) else _W.get(name, default))
    hb=_v("hourly_brutto"); vac=_v("vac_pct"); hol=_v("holiday_pct"); m13=_v("thirteenth_pct")
    exp=_v("expenses_per_hour"); ahv=_v("ahv_pct"); nbu=_v("nbu_pct"); ktg=_v("ktg_pct"); bvg=(p.bvg_pct if p and p.bvg_pct is not None else 0.0)
    gross_hr = hb*(1+(vac+hol+m13)/100.0) + exp
    net_hr   = hb*(1+(vac+hol+m13)/100.0)*(1-(ahv+nbu+ktg+bvg)/100.0) + exp
    return (gross_hr, net_hr)

# ---------------------- Lohnformeln / Helfer ----------------------
def _month_range(year: int, month: int) -> Tuple[date, date]:
    start = date(year, month, 1)
    if month == 12: end = date(year+1, 1, 1) - timedelta(days=1)
    else:           end = date(year, month+1, 1) - timedelta(days=1)
    return (start, end)

def _is_worked(code: str) -> bool:
    return (code or "").upper() == "WORK"

_TYPES_FALLBACK = {
    "de": {"work":"Arbeit","vacation":"Urlaub","sick":"Krank","holiday":"Feiertag","off":"Frei","other":"Sonstiges",
           "appointment":"Arzttermin","hospital":"Spital"},
    "en": {"work":"Work","vacation":"Vacation","sick":"Sick","holiday":"Holiday","off":"Off","other":"Other",
           "appointment":"Appointment","hospital":"Hospital"},
    "sr": {"work":"Rad","vacation":"Odmor","sick":"Bolovanje","holiday":"Praznik","off":"Slobodno","other":"Ostalo",
           "appointment":"Pregled","hospital":"Bolnica"},
}
def _type_label_from_code(code: str, lang: str) -> str:
    if not code: return ""
    if code.startswith("CUSTOM:"): return code.split(":",1)[1] or "Custom"
    key = (code or "").lower()
    try:
        from languages import tr
        txt = tr(f"calendar.entry_types.{key}", lang)
        if txt != f"calendar.entry_types.{key}": return txt
    except Exception: pass
    return _TYPES_FALLBACK.get(lang, _TYPES_FALLBACK["de"]).get(key, code)


# ---------------------- Monatswerte berechnen ----------------------

def compute_month(session: Session, user_id: int, year: int, month: int) -> Dict[str, Any]:
    start, end = _month_range(year, month)
    entries = session.execute(
        select(WorkEntry).where(and_(WorkEntry.user_id == user_id, WorkEntry.date >= start, WorkEntry.date <= end))
    ).scalars().all()

    gross_hr, net_hr = _wage_per_hour(session, user_id)

    hours_work = 0.0
    count = int(len(entries))
    surcharge_gross = 0.0
    surcharge_net   = 0.0

    bd: Dict[str, Dict[str, float]] = {}
    for e in entries:
        key = (e.entry_type or "WORK").upper()
        bd.setdefault(key, {"hours":0.0,"count":0})
        bd[key]["hours"] += float(e.hours or 0.0)
        bd[key]["count"] += 1

        if _is_worked(e.entry_type):
            hours_work += float(e.hours or 0.0)
            if ENABLE_SURCHARGES:
                mins_map = calc_surcharges_for_entry(e.date, e.start_time, e.end_time, e.break_minutes)
                s_g, s_n = surcharge_amounts_from_minutes(mins_map, gross_hr, net_hr)
                surcharge_gross += s_g
                surcharge_net   += s_n

    gross_base = gross_hr * hours_work
    net_base   = net_hr   * hours_work

    return {
        "hours_total": hours_work,
        "entries_count": count,
        "gross_base": gross_base,
        "net_base":   net_base,
        "surcharge_gross": surcharge_gross,
        "surcharge_net":   surcharge_net,
        "gross_total": gross_base + surcharge_gross,
        "net_total":   net_base   + surcharge_net,
        "type_breakdown": bd,
        "start": start, "end": end
    }


# ---------------------- PayrollRun speichern ----------------------

def get_or_create_run(session: Session, user_id: int, year: int, month: int) -> PayrollRun:
    run = session.execute(
        select(PayrollRun).where(and_(PayrollRun.user_id == user_id,
                                      PayrollRun.period_year == year,
                                      PayrollRun.period_month == month))
    ).scalars().one_or_none()
    if run is None:
        run = PayrollRun(user_id=user_id, period_year=year, period_month=month)
        session.add(run); session.flush()
    return run

def recalc_and_save(session: Session, user_id: int, year: int, month: int) -> PayrollRun:
    run = get_or_create_run(session, user_id, year, month)
    if run.is_locked:
        return run

    # Monat berechnen (Basis + Zuschläge)
    data = compute_month(session, user_id, year, month)

    gross_total = float(data["gross_total"])
    net_total   = float(data["net_total"])

    # SICK/ACCIDENT: 80 % von 8 h pro Eintrag zusätzlich vergüten
    bd = data.get("type_breakdown", {}) or {}
    sick_cnt     = int((bd.get("SICK")     or {}).get("count", 0))
    accident_cnt = int((bd.get("ACCIDENT") or {}).get("count", 0))
    sick_like    = sick_cnt + accident_cnt
    if sick_like:
        g_hr, n_hr = _effective_hour_rates(session, user_id)
        gross_total += 0.80 * 8.0 * g_hr * sick_like
        net_total   += 0.80 * 8.0 * n_hr * sick_like

    # PayrollRun schreiben
    run.hours_total   = float(data["hours_total"])
    run.entries_count = int(data["entries_count"])
    run.gross_total   = gross_total
    run.net_total     = net_total
    run.type_breakdown_json = json.dumps(bd, ensure_ascii=False)

    session.commit()
    return run


def lock_run(session: Session, user_id: int, year: int, month: int, by_user_id: int) -> PayrollRun:
    run = recalc_and_save(session, user_id, year, month)
    run.is_locked = True
    run.locked_at = datetime.utcnow()
    run.locked_by_user_id = by_user_id
    session.commit(); return run

def unlock_run(session: Session, user_id: int, year: int, month: int, reason: str) -> PayrollRun:
    run = get_or_create_run(session, user_id, year, month)
    run.is_locked = False
    run.reopen_reason = (reason or "")[:1000] or None
    session.commit(); return run


# ---------------------- PDF Export (Monats-Range, Branding/Logo) ----------------------

def export_pdf(session: Session, user_id: int,
               year_from: int, month_from: int, year_to: int, month_to: int,
               out_path: str, lang: str = "de") -> str:
    # (imports & setup wie in deiner Datei)
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib import colors
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, Image
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.enums import TA_LEFT
        from reportlab.lib.units import cm
        from reportlab.pdfgen import canvas as _canvas
    except ImportError:
        raise RuntimeError("reportlab nicht installiert (pip install reportlab)")

    user = session.query(User).filter(User.id == user_id).one_or_none()
    prof = session.query(UserProfile).filter(UserProfile.user_id == user_id).one_or_none()
    name = f"{(prof.last_name if prof else '')} {(prof.first_name if prof else '')}".strip() or (user.username if user else f"User {user_id}")
    employer = prof.employer if prof and prof.employer else "—"
    employee_id = prof.employee_id if prof and prof.employee_id else ""
    birthday = prof.birthday.strftime("%Y-%m-%d") if (prof and prof.birthday) else ""
    address = prof.address or ""
    postcode = prof.postcode or ""
    city = prof.city or ""
    email = prof.email or ""
    ahv_no = prof.ahv_number if prof and prof.ahv_number else ""  # <— AHV hier sauber gesetzt

    # Monatsrange & period_label (für PDF-Kopf)
    months: List[Tuple[int, int]] = []
    y, m = year_from, month_from
    while (y < year_to) or (y == year_to and m <= month_to):
        months.append((y, m))
        if m == 12: y, m = y+1, 1
        else:       m += 1
    period_label = f"{months[0][0]}-{months[0][1]:02d}" if len(months) == 1 else f"{months[0][0]}-{months[0][1]:02d} bis {months[-1][0]}-{months[-1][1]:02d}"

    # Seitenzahlen Canvas
    class PageNumCanvas(_canvas.Canvas):
        def __init__(self, *args, **kwargs):
            _canvas.Canvas.__init__(self, *args, **kwargs)
            self._saved = []
        def showPage(self): self._saved.append(dict(self.__dict__)); self._startPage()
        def save(self):
            total_pages = len(self._saved)
            for i, state in enumerate(self._saved, start=1):
                self.__dict__.update(state)
                w = self._pagesize[0]
                self.setFont("Helvetica", 9)
                self.drawCentredString(w/2.0, 0.5*cm, f"Seite {i} von {total_pages}")
                _canvas.Canvas.showPage(self)
            _canvas.Canvas.save(self)

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("title_left", parent=styles["Title"], alignment=TA_LEFT)
    line_style  = ParagraphStyle("line_left",  parent=styles["Normal"], alignment=TA_LEFT, fontName="Helvetica", fontSize=10, leading=14)

    # Temp → atomar ersetzen
    out_path = str(out_path)
    out_dir = os.path.dirname(out_path) or "."
    os.makedirs(out_dir, exist_ok=True)
    from uuid import uuid4
    tmp_path = os.path.join(out_dir, f".tmp_payroll_{uuid4().hex}.pdf")

    doc = SimpleDocTemplate(tmp_path, pagesize=A4,
                            topMargin=1*cm, bottomMargin=1*cm, leftMargin=1.5*cm, rightMargin=1.5*cm)
    elems: List[Any] = []

    # Optionales Logo
    logo_path = Path(COMPANY_LOGO_PATH)
    if logo_path.exists():
        try:
            elems.append(Image(str(logo_path), width=3.5*cm, height=3.5*cm))
            elems.append(Spacer(1, 6))
        except Exception:
            pass

    # Titel
    elems.append(Paragraph("Abrechnung", title_style))
    elems.append(Spacer(1, 8))

    # 2-Spalten Kopf linksbündig (6 Zeilen)
    created_str = datetime.now().strftime("%d.%m.%Y")
    header_pairs = [
        ("Name:", name, "Arbeitgeber:", employer),
        ("Geburtsdatum:", birthday, "Mitarbeiter-ID:", employee_id),
        ("Adresse:", address, "AHV-Nr.:", ahv_no),
        ("Postleitzahl:", postcode, "", ""),
        ("Ort:", city, "Zeitraum:", period_label),
        ("E-Mail:", email, "Erstellt am:", created_str),
    ]
    for L, Lv, R, Rv in header_pairs:
        ltxt = f"<b>{L}</b> {Lv}" if L else (Lv or "")
        rtxt = f"<b>{R}</b> {Rv}" if R else (Rv or "")
        row_tbl = Table([[Paragraph(ltxt, line_style), Paragraph(rtxt, line_style)]],
                        colWidths=[260, 260], hAlign='LEFT')
        row_tbl.setStyle(TableStyle([("VALIGN",(0,0),(-1,-1),"TOP")]))
        elems.append(row_tbl)
    elems.append(Spacer(1, 12))

    # Monatsabschnitte
    for (yy, mm) in months:
        elems.append(Paragraph(f"Abrechnung {yy}-{mm:02d}", styles["Heading2"]))
        elems.append(Spacer(1, 6))

        data = compute_month(session, user_id, yy, mm)
        kpis = [
            ["Stunden Total (nur Arbeit)", f"{data['hours_total']:.2f}"],
            ["Einträge", f"{sum(v['count'] for v in data['type_breakdown'].values())}"],
            ["Brutto (Basis)", fmt_money(data['gross_base'])],
            ["Zuschläge (Brutto)", fmt_money(data['surcharge_gross'])],
            ["Brutto (Total)", fmt_money(data['gross_total'])],
            ["Netto (Basis)", fmt_money(data['net_base'])],
            ["Zuschläge (Netto)", fmt_money(data['surcharge_net'])],
            ["Netto (Total)", fmt_money(data['net_total'])],
        ]
        kpi_tbl = Table(kpis, colWidths=[200, 160], hAlign='LEFT')
        kpi_tbl.setStyle(TableStyle([
            ("GRID",(0,0),(-1,-1),0.25,colors.gray),
            ("BACKGROUND",(0,0),(-1,0),colors.HexColor("#eeeeee")),
            ("FONT",(0,0),(-1,0),"Helvetica-Bold"),
            ("ALIGN",(1,1),(1,-1),"RIGHT"),
        ]))
        elems.append(kpi_tbl); elems.append(Spacer(1, 8))

        rows = [["Typ","Stunden","Einträge"]]
        for k, v in data["type_breakdown"].items():
            rows.append([_type_label_from_code(k, "de"),
                         f"{float(v.get('hours',0.0)):.2f}",
                         str(int(v.get('count',0)))])
        br_tbl = Table(rows, repeatRows=1, colWidths=[220, 120, 120], hAlign='LEFT')
        br_tbl.setStyle(TableStyle([
            ("GRID",(0,0),(-1,-1),0.25,colors.gray),
            ("BACKGROUND",(0,0),(-1,0),colors.HexColor("#eeeeee")),
            ("FONT",(0,0),(-1,0),"Helvetica-Bold"),
            ("ALIGN",(1,1),(2,-1),"RIGHT"),
        ]))
        elems.append(br_tbl); elems.append(Spacer(1, 12))

    doc.build(elems, canvasmaker=PageNumCanvas)
    try:
        os.replace(tmp_path, out_path)
    except PermissionError:
        try: os.remove(tmp_path)
        except Exception: pass
        raise RuntimeError("Die Zieldatei kann nicht überschrieben werden (vermutlich im Viewer geöffnet).")
    return out_path


# ---------------------- Excel Export (Multi-Month Range) ----------------------

def export_xlsx_range(session: Session, user_id: int,
                      year_from: int, month_from: int, year_to: int, month_to: int,
                      out_path: str, lang: str = "de") -> str:
    try:
        import openpyxl
        from openpyxl.styles import Font
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise RuntimeError("openpyxl nicht installiert (pip install openpyxl)")

    user = session.query(User).filter(User.id == user_id).one_or_none()
    prof = session.query(UserProfile).filter(UserProfile.user_id == user_id).one_or_none()
    name = f"{(prof.last_name if prof else '')} {(prof.first_name if prof else '')}".strip() or (user.username if user else f"User {user_id}")
    employer = prof.employer if prof and prof.employer else "—"
    employee_id = prof.employee_id if prof and prof.employee_id else ""
    birthday = prof.birthday.strftime("%Y-%m-%d") if (prof and prof.birthday) else ""
    address = prof.address or "";
    postcode = prof.postcode or ""; city = prof.city or ""; email = prof.email or ""
    created_str = datetime.now().strftime("%d.%m.%Y")
    ahv_no = prof.ahv_number if prof and prof.ahv_number else ""  # <— AHV hier gesetzt

    # Range
    months: List[Tuple[int, int]] = []
    y, m = year_from, month_from
    while (y < year_to) or (y == year_to and m <= month_to):
        months.append((y, m))
        if m == 12: y, m = y+1, 1
        else:       m += 1

    wb = openpyxl.Workbook()
    bold = Font(bold=True)

    # Überblick-Sheet (wie gehabt) ...
    ws0 = wb.active; ws0.title = "Überblick"
    ws0.append(["Monat", "Stunden (Arbeit)", "Brutto (Basis)", "Zuschläge Brutto", "Brutto (Total)",
                "Netto (Basis)", "Zuschläge Net", "Netto (Total)", "Einträge"])
    totals = [0.0]*7
    for (yy, mm) in months:
        data = compute_month(session, user_id, yy, mm)
        row = [f"{yy}-{mm:02d}", float(data["hours_total"]), float(data["gross_base"]),
               float(data["surcharge_gross"]), float(data["gross_total"]),
               float(data["net_base"]), float(data["surcharge_net"]),
               float(data["net_total"]), int(sum(v['count'] for v in data["type_breakdown"].values()))]
        ws0.append(row)
        for i in range(7): totals[i] += float(row[i+1])
    ws0.append(["GESAMT", *totals, ""])
    for c in range(1, 10):
        ws0.cell(row=1, column=c).font = bold
        ws0.column_dimensions[get_column_letter(c)].width = 18

    # pro Monat ein eigenes Blatt
    for (yy, mm) in months:
        ws = wb.create_sheet(title=f"{yy}-{mm:02d}")
        month_label = f"{yy}-{mm:02d}"            # <— statt period_label
        head_rows = [
            ("Name:", name, "Arbeitgeber:", employer),
            ("Geburtsdatum:", birthday, "Mitarbeiter-ID:", employee_id),
            ("Adresse:", address, "AHV-Nr.:", ahv_no),
            ("Postleitzahl:", postcode, "", ""),
            ("Ort:", city, "Zeitraum:", month_label),   # <— Monatslabel
            ("E-Mail:", email, "Erstellt am:", created_str),
        ]
        r=1
        for L, Lv, R, Rv in head_rows:
            ws.cell(row=r, column=1, value=L).font = bold
            ws.cell(row=r, column=2, value=Lv)
            ws.cell(row=r, column=3, value=R).font = bold if R else Font(bold=False)
            ws.cell(row=r, column=4, value=Rv); r+=1
        r += 1

        data = compute_month(session, user_id, yy, mm)
        # KPIs + Breakdown wie bisher ...
        kpis = [
            ("Stunden Total (nur Arbeit)", float(data["hours_total"])),
            ("Brutto (Basis)", float(data["gross_base"])),
            ("Zuschläge (Brutto)", float(data["surcharge_gross"])),
            ("Brutto (Total)", float(data["gross_total"])),
            ("Netto (Basis)", float(data["net_base"])),
            ("Zuschläge (Netto)", float(data["surcharge_net"])),
            ("Netto (Total)", float(data["net_total"])),
        ]
        for k,v in kpis:
            ws.cell(row=r, column=1, value=k).font = bold; ws.cell(row=r, column=2, value=v); r+=1
        r += 1

        ws.cell(row=r, column=1, value="Typ").font = bold
        ws.cell(row=r, column=2, value="Stunden").font = bold
        ws.cell(row=r, column=3, value="Einträge").font = bold; r+=1
        for k, v in data["type_breakdown"].items():
            ws.append([_type_label_from_code(k, lang), float(v.get("hours",0.0)), int(v.get("count",0))]); r+=1

        from openpyxl.utils import get_column_letter
        for c in range(1,5):
            ws.column_dimensions[get_column_letter(c)].width = 24

    out_dir = os.path.dirname(str(out_path)) or "."
    os.makedirs(out_dir, exist_ok=True)
    wb.save(str(out_path))
    return str(out_path)
